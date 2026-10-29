# -*- coding: utf-8 -*-
"""Helpers de layout para a apresentacao Lab 9 - OmniDrive."""
import math, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ----------------------------------------------------------------- paleta
BG     = RGBColor(0x0B, 0x12, 0x20)
PANEL  = RGBColor(0x16, 0x22, 0x38)
PANEL2 = RGBColor(0x1E, 0x2E, 0x4A)
TXT    = RGBColor(0xE6, 0xED, 0xF8)
MUTED  = RGBColor(0x93, 0xA6, 0xC4)
DIM    = RGBColor(0x5A, 0x6C, 0x8A)
CY     = RGBColor(0x38, 0xD9, 0xF0)
AM     = RGBColor(0xFB, 0xBF, 0x24)
GR     = RGBColor(0x34, 0xD3, 0x99)
PK     = RGBColor(0xF4, 0x72, 0xB6)
VI     = RGBColor(0xA7, 0x8B, 0xFA)
OR     = RGBColor(0xFB, 0x8E, 0x3C)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CMT    = RGBColor(0x6B, 0x82, 0xA6)

UI   = "Segoe UI"
MONO = "Consolas"

W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)

STATE = {"sec": "", "n": 0}


# ----------------------------------------------------------------- primitivas
def tb(slide, x, y, w, h, text, size=18, color=TXT, bold=False, italic=False,
       align=PP_ALIGN.LEFT, font=UI, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    w, h = max(0.02, w), max(0.02, h)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=PANEL, line=None, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None, rot=None):
    w, h = max(0.02, w), max(0.02, h)
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    if rot is not None:
        sh.rotation = rot
    if sh.has_text_frame:
        sh.text_frame.word_wrap = True
    return sh


def shape_text(sh, text, size=14, color=TXT, bold=False, font=UI,
               align=PP_ALIGN.CENTER, spacing=1.05):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, ln in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name, f.size, f.bold = font, Pt(size), bold
        f.color.rgb = color
    return sh


def line(slide, x1, y1, x2, y2, color=MUTED, lw=1.25, dash=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    if dash:
        cn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return cn


def arrow_r(slide, x, y, w, h=0.28, color=CY):
    return rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def arrow_d(slide, x, y, h, w=0.28, color=CY):
    return rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.DOWN_ARROW)


def arrow_l(slide, x, y, w, h=0.28, color=CY):
    return rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.LEFT_ARROW)


def vec(slide, x1, y1, x2, y2, color=CY, lw=2.25, head=0.24):
    """Vetor: haste + ponta triangular orientada."""
    ang = math.degrees(math.atan2(y2 - y1, x2 - x1))
    d = math.hypot(x2 - x1, y2 - y1)
    if d > head:
        bx = x1 + (x2 - x1) * (d - head * 0.7) / d
        by = y1 + (y2 - y1) * (d - head * 0.7) / d
    else:
        bx, by = x1, y1
    line(slide, x1, y1, bx, by, color=color, lw=lw)
    rect(slide, x2 - head / 2, y2 - head / 2, head, head, fill=color,
         shape=MSO_SHAPE.ISOSCELES_TRIANGLE, rot=ang + 90)


# ----------------------------------------------------------------- slide base
def new_slide(kicker=None, title=None, sec=None, number=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, -0.1, -0.1, W + 0.2, H + 0.2, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    if sec is not None:
        STATE["sec"] = sec
    if kicker is not None:
        tb(s, 0.75, 0.42, 10.0, 0.3, kicker.upper(), size=12, color=CY, bold=True)
    if title is not None:
        tb(s, 0.75, 0.76, 11.9, 0.95, title, size=30, color=WHITE, bold=True, spacing=1.0)
        rect(s, 0.75, 1.76, 1.5, 0.045, fill=AM, shape=MSO_SHAPE.RECTANGLE)
    if number:
        STATE["n"] += 1
        tb(s, 0.75, H - 0.52, 8.0, 0.3, STATE["sec"], size=10.5, color=DIM)
        tb(s, W - 1.6, H - 0.52, 0.85, 0.3, "%02d" % STATE["n"], size=10.5,
           color=DIM, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, x, y, w, items, size=16.5, gap=0.1, lead_color=AM, spacing=1.18):
    """items: str | (lead, resto) | ("sub", texto)"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = spacing
        p.space_after = Pt(gap * 72)
        if isinstance(it, tuple) and it[0] == "sub":
            r = p.add_run()
            r.text = u"      –  " + it[1]
            r.font.name, r.font.size = UI, Pt(size - 1.5)
            r.font.color.rgb = MUTED
            continue
        r = p.add_run()
        r.text = u"▸  "
        r.font.name, r.font.size = UI, Pt(size)
        r.font.color.rgb = CY
        if isinstance(it, tuple):
            a = p.add_run()
            a.text = it[0]
            a.font.name, a.font.size, a.font.bold = UI, Pt(size), True
            a.font.color.rgb = lead_color
            b = p.add_run()
            b.text = it[1]
            b.font.name, b.font.size = UI, Pt(size)
            b.font.color.rgb = TXT
        else:
            a = p.add_run()
            a.text = it
            a.font.name, a.font.size = UI, Pt(size)
            a.font.color.rgb = TXT
    return box


# ----------------------------------------------------------------- codigo
JAVA_KW = set("""public private protected static final class void double int long boolean
new return if else for while this extends implements import package true false null super
abstract interface throws""".split())
TOKEN = re.compile(r'(//[^\n]*|"[^"]*"|@?\b[A-Za-z_][A-Za-z0-9_]*\b|\d+\.?\d*|\s+|.)')


def code(slide, x, y, w, h, src, size=12.5, title=None):
    lines = src.split("\n")
    nl = max(1, len(lines))
    maxlen = max(1, max(len(l) for l in lines))
    # Consolas: avanco ~0.552 em. Encolhe para caber na largura E na altura.
    s_h = 72.0 * (w - 0.42) / (0.552 * maxlen)
    s_v = 72.0 * (h - 0.34) / (nl * 1.45)
    size = max(8.5, min(size, s_h, s_v))
    box = rect(slide, x, y, w, h, fill=PANEL, line=PANEL2, lw=1.0, adj=0.03)
    if title:
        tw = min(w - 0.3, 0.3 + 0.083 * len(title))
        tag = rect(slide, x + 0.18, y - 0.17, tw, 0.33, fill=PANEL2, adj=0.4)
        shape_text(tag, title, size=10.5, color=CY, bold=True, font=MONO)
    tf = box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.24)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.1)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.2
        for tok in TOKEN.findall(ln):
            r = p.add_run()
            r.text = tok.replace("\t", "    ")
            f = r.font
            f.name, f.size = MONO, Pt(size)
            if tok.startswith("//"):
                f.color.rgb = CMT
                f.italic = True
            elif tok.startswith('"'):
                f.color.rgb = AM
            elif tok in JAVA_KW:
                f.color.rgb = VI
            elif re.fullmatch(r'\d+\.?\d*', tok):
                f.color.rgb = PK
            elif tok.startswith("@"):
                f.color.rgb = OR
            elif re.fullmatch(r'[A-Z][A-Za-z0-9_]*', tok):
                f.color.rgb = GR
            else:
                f.color.rgb = TXT
    return box


def card(slide, x, y, w, h, head, body, accent=CY, hsize=15, bsize=13):
    rect(slide, x, y, w, h, fill=PANEL, line=PANEL2, adj=0.08)
    rect(slide, x, y + 0.14, 0.055, h - 0.28, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    tb(slide, x + 0.3, y + 0.2, w - 0.52, 0.35, head, size=hsize, bold=True, color=accent)
    tb(slide, x + 0.3, y + 0.58, w - 0.52, h - 0.72, body, size=bsize, color=TXT, spacing=1.2)


def formula(slide, x, y, w, h, text, size=19, color=WHITE, fill=PANEL2,
            align=PP_ALIGN.CENTER, font=MONO):
    sh = rect(slide, x, y, w, h, fill=fill, line=None, adj=0.1)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.16)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.32
        r = p.add_run()
        r.text = ln
        r.font.name, r.font.size = font, Pt(size)
        r.font.color.rgb = color
    return sh


def table(slide, x, y, w, rows, colw, size=13, rowh=0.44, mono_body=True):
    for ri, row in enumerate(rows):
        yy = y + ri * rowh
        if ri == 0:
            rect(slide, x, yy, w, rowh, fill=PANEL2, shape=MSO_SHAPE.RECTANGLE)
        elif ri % 2 == 0:
            rect(slide, x, yy, w, rowh, fill=PANEL, shape=MSO_SHAPE.RECTANGLE)
        cx = x
        for ci, cell in enumerate(row):
            cw = colw[ci]
            tb(slide, cx + 0.16, yy, cw - 0.24, rowh, str(cell),
               size=size if ri else size - 0.5,
               color=CY if ri == 0 else TXT, bold=(ri == 0),
               font=UI if (ri == 0 or not mono_body) else MONO,
               anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
            cx += cw
    return y + len(rows) * rowh


def note(slide, x, y, w, text, color=AM, size=13.5, icon=u"⚠", h=None):
    if h is None:
        h = max(0.56, 0.34 + 0.24 * math.ceil(len(text) / 100.0))
    if y + h > H - 0.60:                      # nunca encosta no rodape
        y = H - 0.60 - h
    bgc = RGBColor(0x26, 0x20, 0x10) if color is AM else PANEL
    rect(slide, x, y, w, h, fill=bgc, line=color, lw=1.0, adj=0.1)
    tb(slide, x + 0.22, y + (h - 0.3) / 2, 0.34, 0.3, icon, size=size + 2, color=color, bold=True)
    tb(slide, x + 0.66, y + 0.1, w - 0.9, h - 0.2, text, size=size, color=TXT,
       spacing=1.18, anchor=MSO_ANCHOR.MIDDLE)


def section(num, title, subtitle, accent=CY, topics=()):
    s = new_slide(sec=title, number=False)
    rect(s, 0, 0, 0.22, H, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 9.6, 1.45, 3.2, 3.4, num, size=190, color=PANEL2, bold=True,
       align=PP_ALIGN.CENTER)
    tb(s, 1.45, 2.62, 9.0, 0.4, "PARTE " + num, size=13, color=accent, bold=True)
    tb(s, 1.45, 3.0, 10.4, 0.95, title, size=40, color=WHITE, bold=True)
    tb(s, 1.45, 4.05, 9.8, 0.6, subtitle, size=17, color=MUTED)
    if topics:
        xx = 1.45
        for t in topics:
            wd = 0.38 + 0.097 * len(t)
            sh = rect(s, xx, 4.95, wd, 0.4, fill=PANEL, line=PANEL2, adj=0.45)
            shape_text(sh, t, size=11.5, color=MUTED)
            xx += wd + 0.14
    return s
